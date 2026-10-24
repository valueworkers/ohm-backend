from django.db import connection, transaction
from django.db.models import Q
from django.db.models.deletion import ProtectedError
from django.shortcuts import get_object_or_404
from django.core import signing
from django.utils import timezone
from rest_framework.decorators import action
from rest_framework import mixins, status, viewsets
from rest_framework.exceptions import AuthenticationFailed,ValidationError
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework_simplejwt.tokens import RefreshToken

from django.contrib.auth.models import Permission
from apps.accounts.models import (
    EmployeeProfile, Role, RoleAssignment, ShiftSchedule, StaffForHire,
    User, UserDocument, UserDocumentFile,
)
from apps.accounts.permissions import IsTenantAdmin, IsTenantSchema
from apps.accounts.serializers import (
    CustomerRegisterSerializer,
    BulkEmployeeUploadSerializer,
    OTPPasswordResetSerializer,
    OTPRequestSerializer,
    UnifiedLoginSerializer,
    OwnerCreateSerializer,
    OwnerSerializer,
    UserSerializer,
    RoleSerializer,
    RoleAssignmentSerializer,
    CustomerListSerializer,
    EmployeeTerminationSerializer,
    MedicareEmployeeSerializer,
    MedicareUserSerializer,
    ShiftScheduleSerializer,
    StaffForHireSerializer,
    UserDocumentSerializer,
)

from .helpers import consume_otp, find_user, has_password, issue_otp

# ------------------------ Shared response helpers ------------------------
def _invalid(message="Invalid credentials."):
    return Response({"error": message}, status=status.HTTP_401_UNAUTHORIZED)

# ------------------------ Authentication: login / logout ------------------------
class UnifiedLoginView(APIView):
    """Login with email or phone, using either password or OTP."""

    authentication_classes = []
    permission_classes = [AllowAny]
    # throttle_classes = [ScopedRateThrottle]
    # throttle_scope = "root_login"

    def post(self, request):
        serializer = UnifiedLoginSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        user = find_user(data["username"])
        if not user:
            return _invalid() if data["method"] == "password" else _invalid("Invalid or expired OTP.")

        # --- Verify credentials (password or OTP) ---
        if data["method"] == "password":
            if not has_password(user):
                if user.is_active:
                    return Response(
                        {
                            "error": "Password has not been set for this account. Please try OTP login.",
                            "code": "password_not_set",
                        },
                        status=status.HTTP_400_BAD_REQUEST,
                    )
                return _invalid()
            if not user.check_password(data["password"]):
                return _invalid()
        else:
            if not consume_otp(user, "login", data["otp"]):
                return _invalid("Invalid or expired OTP.")

        # --- Account state checks (credentials already verified) ---
        if not user.is_active:
            return Response(
                {"error": "Account pending approval by Master Admin."},
                status=status.HTTP_403_FORBIDDEN,
            )

        # --- Issue tokens ---
        user.last_login = timezone.now()
        user.save(update_fields=["last_login"])

        refresh = RefreshToken.for_user(user)
        schema_name = connection.schema_name
        refresh["schema"] = schema_name
        return Response(
            {
                "message": "Login successful!",
                "schema_name": schema_name,
                "user": MedicareUserSerializer(user, context={"request": request}).data,
                "tokens": {"refresh": str(refresh), "access": str(refresh.access_token)},
            },
            status=status.HTTP_200_OK,
        )


class UnifiedLogoutView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        try:
            refresh_token = request.data.get("refresh_token")
            if refresh_token:
                RefreshToken(refresh_token).blacklist()
            return Response({"message": "Logout successful"}, status=status.HTTP_200_OK)
        except Exception:
            return Response({"error": "Invalid token or logout failed"}, status=status.HTTP_400_BAD_REQUEST)


class OwnerViewSet(viewsets.ModelViewSet):
    """
    /owners/            GET list, POST add co-owner
    /owners/{id}/       GET, PATCH/PUT, DELETE (soft: deactivates)
    /owners/me/         GET, PATCH own profile
    """

    http_method_names = ["get", "post", "patch", "put", "delete", "head", "options"]

    def owners_qs(self):
        return User.objects.filter(
            role_assignments__role__code=Role.Code.TENANT_ADMIN,
            role_assignments__scope_type__isnull=True,
        ).distinct()

    def get_queryset(self):
        return self.owners_qs().order_by("date_joined")

    def get_serializer_class(self):
        return OwnerCreateSerializer if self.action == "create" else OwnerSerializer

    def create(self, request, *args, **kwargs):
        ser = self.get_serializer(data=request.data)
        ser.is_valid(raise_exception=True)
        user = ser.save()
        return Response(OwnerSerializer(user).data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=["get", "patch"], url_path="me")
    def me(self, request):
        if request.method == "GET":
            return Response(OwnerSerializer(request.user).data)
        ser = OwnerSerializer(request.user, data=request.data, partial=True)
        ser.is_valid(raise_exception=True)
        ser.save()
        return Response(ser.data)

    def destroy(self, request, *args, **kwargs):
        owner = self.get_object()
        if owner.pk == request.user.pk:
            return Response({"detail": "You cannot remove yourself."}, status=400)
        if not self.owners_qs().filter(is_active=True).exclude(pk=owner.pk).exists():
            return Response({"detail": "Cannot remove the last owner."}, status=400)
        owner.is_active = False
        owner.save(update_fields=["is_active"])
        return Response(status=status.HTTP_204_NO_CONTENT)

# ------------------------ Authentication: OTP request / password reset ------------------------
class OTPRequestView(APIView):
    authentication_classes = []
    permission_classes = [AllowAny]
    # throttle_classes = [ScopedRateThrottle]
    # throttle_scope = "otp_request"

    def post(self, request):
        serializer = OTPRequestSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        user = find_user(data["username"])
        # Identical response for unknown, inactive and ineligible accounts.
        if user and user.is_active:
            issue_otp(user, data["username"], data["purpose"])
        return Response(
            {"detail": "If the account exists, a verification code has been sent."},
            status=status.HTTP_202_ACCEPTED,
        )


class OTPPasswordResetView(APIView):
    authentication_classes = []
    permission_classes = [AllowAny]
    # throttle_classes = [ScopedRateThrottle]
    # throttle_scope = "otp_request"

    def post(self, request):
        serializer = OTPPasswordResetSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        user = find_user(data["username"])
        if not user or not user.is_active:
            raise AuthenticationFailed("Invalid or expired OTP.")
        if not consume_otp(user, "password_reset", data["otp"]):
            raise AuthenticationFailed("Invalid or expired OTP.")

        user.set_password(data["new_password"])
        user.save(update_fields=("password",))
        return Response({"detail": "Password has been reset."})

# ------------------------ Tenant: customer self-registration ------------------------
class TenantRegistrationCreateView(APIView):
    """Let a provider's customers create login accounts in that provider's schema."""

    authentication_classes = []
    permission_classes = [AllowAny, IsTenantSchema]
    # throttle_classes = [ScopedRateThrottle]
    # throttle_scope = "tenant_registration_create"

    def post(self, request):
        serializer = CustomerRegisterSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        with transaction.atomic():
            user = User.objects.create_user(
                **data,
            )
            RoleAssignment.objects.create(user=user, role=Role.objects.get(code=Role.Code.CUSTOMER))
        return Response({"email": user.email}, status=status.HTTP_201_CREATED)

# ------------------------ Tenant: user management (Tenant Admin only) ------------------------
class UserViewSet(
    mixins.ListModelMixin,
    mixins.CreateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.UpdateModelMixin,
    viewsets.GenericViewSet,
):
    """Let a provider's Tenant Admin manage accounts stored in that provider's schema."""

    serializer_class = UserSerializer
    http_method_names = ["get", "post", "patch", "head", "options"]

    def get_queryset(self):
        queryset = User.objects.all().order_by("id").prefetch_related("role_assignments__role")

        role = self.request.query_params.get("role")
        if role:
            queryset = queryset.filter(
                role_assignments__role__code=role,
                role_assignments__scope_type__isnull=True,
            )

        search = self.request.query_params.get("q")
        if search:
            queryset = queryset.filter(
                Q(email__icontains=search) | Q(first_name__icontains=search)
                | Q(middle_name__icontains=search) | Q(last_name__icontains=search)
            )

        return queryset.distinct()


class RoleViewSet(viewsets.ModelViewSet):
    """Tenant Admin CRUD for workspace roles; core roles remain protected."""

    serializer_class = RoleSerializer
    queryset = Role.objects.prefetch_related("permissions").order_by("rank", "name")
    http_method_names = ["get", "post", "patch", "put", "delete", "head", "options"]

    @action(detail=False, methods=["get"])
    def permissions(self, request):
        permissions = Permission.objects.order_by("codename")
        return Response([
            {"id": permission.id, "code": permission.codename, "name": permission.name}
            for permission in permissions
        ])

    def destroy(self, request, *args, **kwargs):
        role = self.get_object()
        if role.code in {Role.Code.TENANT_ADMIN, Role.Code.CUSTOMER}:
            raise ValidationError({"code": "Built-in roles cannot be deleted."})
        try:
            return super().destroy(request, *args, **kwargs)
        except ProtectedError:
            return Response(
                {"detail": "This role is assigned to users. Remove its assignments before deleting it."},
                status=status.HTTP_409_CONFLICT,
            )


class RoleAssignmentViewSet(
    mixins.ListModelMixin,
    mixins.CreateModelMixin,
    mixins.DestroyModelMixin,
    viewsets.GenericViewSet,
):
    """Tenant Admin role grants, optionally scoped to an object in this tenant."""

    serializer_class = RoleAssignmentSerializer
    queryset = RoleAssignment.objects.select_related("user", "role", "scope_type").order_by("id")
    http_method_names = ["get", "post", "delete", "head", "options"]

    def get_queryset(self):
        queryset = super().get_queryset()
        user_id = self.request.query_params.get("user_id")
        role_id = self.request.query_params.get("role_id")
        if user_id:
            queryset = queryset.filter(user_id=user_id)
        if role_id:
            queryset = queryset.filter(role_id=role_id)
        return queryset

    def perform_destroy(self, instance):
        if (instance.scope_type_id is None and instance.role.code == Role.Code.TENANT_ADMIN
                and instance.user.is_active):
            another_admin = User.objects.filter(
                is_active=True,
                role_assignments__role__code=Role.Code.TENANT_ADMIN,
                role_assignments__scope_type__isnull=True,
            ).exclude(pk=instance.user_id).exists()
            if not another_admin:
                raise ValidationError({"detail": "A workspace must keep at least one active Tenant Admin."})
        instance.delete()


class ShiftScheduleViewSet(viewsets.ModelViewSet):
    queryset = ShiftSchedule.objects.all()
    serializer_class = ShiftScheduleSerializer
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    http_method_names = ["get", "post", "put", "patch", "delete", "head", "options"]


class EmployeeViewSet(viewsets.ModelViewSet):
    serializer_class = MedicareEmployeeSerializer
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    search_fields = ["first_name", "middle_name", "last_name", "email", "mobile_number", "employee_profile__employee_id"]
    ordering_fields = ["id", "first_name", "last_name", "date_joined"]
    ordering = ["id"]

    def get_queryset(self):
        queryset = User.objects.employees().select_related("employee_profile", "employee_profile__shift")
        if self.action == "list":
            term = self.request.query_params.get("q")
            if term:
                queryset = queryset.filter(Q(email__icontains=term) | Q(first_name__icontains=term) | Q(last_name__icontains=term) | Q(employee_profile__employee_id__icontains=term))
            active = self.request.query_params.get("is_active")
            if active is not None:
                queryset = queryset.filter(is_active=active.lower() == "true")
            category = self.request.query_params.get("category")
            if category:
                queryset = queryset.filter(employee_profile__category=category)
            termination = self.request.query_params.get("termination_type")
            if termination:
                queryset = queryset.filter(employee_profile__termination_type=termination)
        return queryset.order_by(*self.ordering)

    def destroy(self, request, *args, **kwargs):
        employee = self.get_object()
        profile = get_object_or_404(EmployeeProfile, user=employee)
        serializer = EmployeeTerminationSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        try:
            profile.terminate(data["termination_type"], data.get("reason"), data.get("last_working_day"))
        except Exception as exc:
            if hasattr(exc, "message_dict"):
                raise ValidationError(exc.message_dict)
            raise
        return Response({"message": "Employee terminated."})

    @action(detail=True, methods=["post"], url_path="terminate")
    def terminate_employee(self, request, pk=None):
        employee = self.get_object()
        profile = get_object_or_404(EmployeeProfile, user=employee)
        serializer = EmployeeTerminationSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        try:
            profile.terminate(data["termination_type"], data.get("reason"), data.get("last_working_day"))
        except Exception as exc:
            if hasattr(exc, "message_dict"):
                raise ValidationError(exc.message_dict)
            raise
        return Response(MedicareEmployeeSerializer(employee, context=self.get_serializer_context()).data)

    @action(detail=True, methods=["post"], url_path="termination-revoke")
    def termination_revoke(self, request, pk=None):
        employee = self.get_object()
        profile = get_object_or_404(EmployeeProfile, user=employee)
        raw = request.data.get("rehired", request.data.get("rehired_status", False))
        rehired = raw is True or str(raw).lower() in ("true", "1", "yes")
        try:
            profile.termination_revoke(rehired)
        except Exception as exc:
            if hasattr(exc, "message_dict"):
                raise ValidationError(exc.message_dict)
            raise
        return Response(MedicareEmployeeSerializer(employee, context=self.get_serializer_context()).data)


class CustomerViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    search_fields = ["first_name", "middle_name", "last_name", "email", "mobile_number", "city"]
    ordering = ["first_name"]

    def get_queryset(self):
        return User.objects.customers().filter(is_deleted=False).order_by("-id")

    def get_serializer_class(self):
        return CustomerListSerializer if self.action == "list" else MedicareUserSerializer

    def perform_create(self, serializer):
        user = serializer.save()
        role = Role.objects.filter(code=Role.Code.CUSTOMER).first()
        if role:
            RoleAssignment.objects.get_or_create(user=user, role=role)

    def destroy(self, request, *args, **kwargs):
        user = self.get_object()
        user.soft_delete()
        return Response({"message": "Customer deleted successfully."})


class UserDocumentViewSet(viewsets.ModelViewSet):
    serializer_class = UserDocumentSerializer
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    parser_classes = [MultiPartParser, FormParser]
    queryset = UserDocument.objects.select_related("user", "uploaded_by").prefetch_related("files")

    def perform_create(self, serializer):
        document = serializer.save(uploaded_by=self.request.user)
        files = self.request.FILES.getlist("files")
        UserDocumentFile.objects.bulk_create([UserDocumentFile(document=document, file=f) for f in files])

    def perform_update(self, serializer):
        document = serializer.save()
        files = self.request.FILES.getlist("files")
        UserDocumentFile.objects.bulk_create([UserDocumentFile(document=document, file=f) for f in files])


class StaffForHireViewSet(viewsets.ModelViewSet):
    serializer_class = StaffForHireSerializer
    permission_classes = [IsAuthenticated]
    queryset = StaffForHire.objects.select_related("created_by")
    search_fields = ["staff_name", "vendor_name", "available_for", "language", "skill"]
    ordering = ["-created_at"]

    def get_permissions(self):
        if self.action in {"list", "retrieve"}:
            return [IsAuthenticated()]
        return [IsAuthenticated(), IsTenantAdmin()]

    def get_queryset(self):
        queryset = super().get_queryset()
        service = self.request.query_params.get("available_for")
        if service:
            queryset = queryset.filter(available_for__contains=[service])
        if not self.request.user.is_owner:
            queryset = queryset.filter(is_active=True, is_available=True)
        return queryset


class UserProfileView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response(MedicareUserSerializer(request.user, context={"request": request}).data)

    def put(self, request):
        serializer = MedicareUserSerializer(request.user, data=request.data, partial=True, context={"request": request})
        for field in ("email", "is_active", "is_deleted", "created_by", "date_joined"):
            if field in serializer.fields:
                serializer.fields[field].read_only = True
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data)


class ChangePasswordView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        current = request.data.get("old_password")
        new = request.data.get("new_password")
        if not current or not new:
            raise ValidationError({"detail": "old_password and new_password are required."})
        if not request.user.check_password(current):
            raise ValidationError({"old_password": "Current password is incorrect."})
        from django.contrib.auth.password_validation import validate_password
        from django.core.exceptions import ValidationError as DjangoValidationError
        try:
            validate_password(new, request.user)
        except DjangoValidationError as exc:
            raise ValidationError({"new_password": exc.messages})
        request.user.set_password(new)
        request.user.save(update_fields=["password"])
        return Response({"message": "Password changed successfully."})


class VerifyOTPView(APIView):
    permission_classes = [AllowAny]
    authentication_classes = []

    def post(self, request):
        username = request.data.get("username")
        otp = request.data.get("otp")
        user = find_user(username or "")
        if not user or not user.is_active or not consume_otp(user, "password_reset", otp or ""):
            raise AuthenticationFailed("Invalid or expired OTP.")
        token = signing.TimestampSigner(salt="accounts.password-reset").sign(f"{connection.schema_name}:{user.pk}")
        return Response({"reset_token": token})


class ResetPasswordView(APIView):
    permission_classes = [AllowAny]
    authentication_classes = []

    def post(self, request):
        token = request.data.get("reset_token")
        password = request.data.get("new_password")
        if not token or not password:
            raise ValidationError({"detail": "reset_token and new_password are required."})
        try:
            from django.contrib.auth.password_validation import validate_password
            from django.core.exceptions import ValidationError as DjangoValidationError
            identity = signing.TimestampSigner(salt="accounts.password-reset").unsign(token, max_age=600)
            schema, user_id = identity.split(":", 1)
            if schema != connection.schema_name:
                raise signing.BadSignature("Token belongs to a different schema.")
            user = User.objects.get(pk=user_id, is_active=True)
            validate_password(password, user)
        except signing.SignatureExpired:
            raise AuthenticationFailed("Password reset token has expired.")
        except (signing.BadSignature, User.DoesNotExist, ValueError):
            raise AuthenticationFailed("Invalid password reset token.")
        except DjangoValidationError as exc:
            raise ValidationError({"new_password": exc.messages})
        user.set_password(password)
        user.save(update_fields=["password"])
        return Response({"message": "Password reset successful."})


class BulkEmployeeUploadAPIView(APIView):
    permission_classes = [IsAuthenticated, IsTenantAdmin]
    parser_classes = [MultiPartParser, FormParser]
    MAX_ROWS = 1000

    def post(self, request):
        serializer = BulkEmployeeUploadSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(serializer.validated_data["file"], read_only=True, data_only=True)
        except Exception as exc:
            raise ValidationError({"file": f"Could not read workbook: {exc}"})
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip().lower().replace(" ", "_") for value in next(rows, ())]
        if not headers or "email" not in headers:
            raise ValidationError({"file": "The first row must contain an email column."})
        results = {"created": [], "failed": []}
        for row_number, row in enumerate(rows, start=2):
            if row_number > self.MAX_ROWS + 1:
                raise ValidationError({"file": f"A maximum of {self.MAX_ROWS} employees can be uploaded at once."})
            record = {header: value for header, value in zip(headers, row) if header}
            if not any(value is not None for value in record.values()):
                continue
            email = str(record.get("email") or "").strip().lower()
            try:
                if not email:
                    raise ValueError("email is required")
                with transaction.atomic():
                    if User.objects.filter(email__iexact=email).exists():
                        raise ValueError("email already exists")
                    profile_fields = {key: record[key] for key in (
                        "employee_id", "category", "designation", "grade", "cost_center", "department",
                        "vendor_name", "vendor_phone", "pf_number", "uan_number", "esi_number", "esi_dispensary",
                    ) if key in record and record[key] is not None}
                    first = str(record.get("first_name") or "")
                    last = str(record.get("last_name") or "")
                    user = User.objects.create_user(
                        email=email, password=None, first_name=first, last_name=last,
                        mobile_number=str(record.get("mobile_number") or "") or None,
                        created_by=request.user,
                    )
                    EmployeeProfile.objects.create(user=user, **profile_fields)
                    code = str(record.get("role") or Role.Code.STAFF).upper()
                    allowed_roles = {Role.Code.STAFF, Role.Code.OPS_MANAGER, Role.Code.VENUE_ADMIN, Role.Code.SERVICE_ADMIN, Role.Code.RESOURCE_ADMIN}
                    if code not in allowed_roles:
                        raise ValueError("role must be a staff or manager role")
                    role = Role.objects.get(code=code)
                    RoleAssignment.objects.get_or_create(user=user, role=role)
                results["created"].append({"row": row_number, "email": email, "id": user.pk})
            except Exception as exc:
                results["failed"].append({"row": row_number, "email": email, "error": str(exc)})
        workbook.close()
        results["summary"] = {"created": len(results["created"]), "failed": len(results["failed"])}
        return Response(results, status=status.HTTP_207_MULTI_STATUS if results["failed"] else status.HTTP_200_OK)
